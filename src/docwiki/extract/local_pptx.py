"""PPTX via python-pptx: slide titles, text, tables, pictures and speaker notes (local adapter).

Every block of slide N carries a ``pptx-slide`` locator (``slide-NN`` sidecar); the optional
``shape`` field names the shape (or ``notes`` for the speaker notes).
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from docwiki.contracts.locators import PptxSlideLocator
from docwiki.extract._draft import Draft
from docwiki.extract._images import embedded_png
from docwiki.extract._tables import emit_table
from docwiki.extract._text import clean_text, lang_from_tag, one_line, safe_label

__all__ = ["extract_pptx"]


def _text(value: str) -> str:
    return clean_text(value.replace("\v", "\n"))


def _flat_shapes(shapes: Any) -> list[Any]:
    flat: list[Any] = []

    def walk(collection: Any) -> None:
        for shape in collection:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes)
            else:
                flat.append(shape)

    walk(shapes)
    return sorted(flat, key=lambda shape: (shape.top or 0, shape.left or 0))


def _grid(table: Any) -> list[list[str]]:
    rows, columns = len(table.rows), len(table.columns)
    grid = [[_text(table.cell(r, c).text) for c in range(columns)] for r in range(rows)]
    for r in range(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            if not cell.is_merge_origin:
                continue
            for rr in range(r, min(rows, r + cell.span_height)):
                for cc in range(c, min(columns, c + cell.span_width)):
                    grid[rr][cc] = grid[r][c]
    return grid


def _picture_blob(shape: Any) -> bytes | None:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE and not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.image.blob
    except (AttributeError, KeyError, ValueError):
        return None


def extract_pptx(data: bytes, draft: Draft, *, max_pages: int) -> None:
    presentation = Presentation(io.BytesIO(data))
    properties = presentation.core_properties
    created = properties.created
    draft.authored = created.date() if created is not None else None
    draft.lang = lang_from_tag(properties.language)
    slides = list(presentation.slides)
    draft.page_count = len(slides)
    if len(slides) > max_pages:
        draft.warn("max-pages", f"extracted the first {max_pages} of {len(slides)} slides")
    first_title: str | None = None
    for number, slide in enumerate(slides[:max_pages], 1):
        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None
        if title_shape is not None and title_shape.has_text_frame:
            title = _text(title_shape.text_frame.text)
            if title:
                first_title = first_title or one_line(title)
                draft.block(
                    "heading",
                    title,
                    PptxSlideLocator(slide=number, shape=_shape_label(title_shape)),
                    "native-text",
                    level=1 if number == 1 else 2,
                )
        for shape in _flat_shapes(slide.shapes):
            if shape.shape_id == title_id:
                continue
            _extract_shape(shape, number, draft)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = _text(frame.text) if frame is not None else ""
            if notes:
                locator = PptxSlideLocator(slide=number, shape="notes")
                draft.block("note", notes, locator, "native-text")
    draft.title_hint = one_line(properties.title or "") or first_title


def _shape_label(shape: Any) -> str:
    return safe_label(shape.name, fallback=f"shape {shape.shape_id}", max_len=60)


def _extract_shape(shape: Any, number: int, draft: Draft) -> None:
    label = _shape_label(shape)
    locator = PptxSlideLocator(slide=number, shape=label)
    if getattr(shape, "has_table", False) and shape.has_table:
        grid = _grid(shape.table)
        if grid:
            # Native text, not deterministic-table: a slide is one sidecar (``slide-NN``) with a
            # single provenance, and slide text is native text.
            emit_table(
                draft,
                header=grid[0],
                rows=[(index, cells) for index, cells in enumerate(grid[1:], 2)],
                locator=locator,
                provenance="native-text",
                row_locator=lambda _row: locator,
                what=f"slide {number}, {label}",
            )
        return
    if shape.has_text_frame:
        text = _text(shape.text_frame.text)
        if text:
            draft.block("slide", text, locator, "native-text")
        return
    blob = _picture_blob(shape)
    if blob is None:
        return
    png = embedded_png(blob)
    if png is None:
        draft.warn("figure-skipped", f"slide {number}: picture format not supported or too small")
        return
    draft.figure(locator, png)
