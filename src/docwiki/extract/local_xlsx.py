"""XLSX via openpyxl: one deterministic table per sheet (values, not formulas).

Row numbers are Excel row numbers and row 1 is the header. Merged ranges are expanded (every
cell of the range carries the top-left value). Formulas contribute their cached value only.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, time, timedelta
from typing import Any

from openpyxl import load_workbook

from docwiki.contracts.ids import assign_unique_slugs, slugify_ascii
from docwiki.contracts.locators import XlsxRowsLocator
from docwiki.extract._draft import Draft
from docwiki.extract._tables import MAX_TABLE_COLUMNS, MAX_TABLE_ROWS, emit_table
from docwiki.extract._text import clean_text, lang_from_tag, one_line, safe_label

__all__ = ["extract_xlsx", "format_cell"]

_DECIMALS_RE = re.compile(r"0\.(0+)")


def format_cell(value: Any, number_format: str | None = None) -> str:
    """Deterministic cell text: ISO dates, percentages as shown, 15 significant digits."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        fmt = number_format or ""
        if "%" in fmt:
            match = _DECIMALS_RE.search(fmt.split(";", 1)[0])
            decimals = len(match[1]) if match else 0
            return f"{value * 100:.{decimals}f}%"
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return format(value, ".15g")
    return clean_text(str(value))


def _sheet_cells(sheet: Any) -> dict[tuple[int, int], Any]:
    cells = getattr(sheet, "_cells", None)
    if isinstance(cells, dict):
        return cells
    return {  # pragma: no cover - openpyxl without the private cell map
        (cell.row, cell.column): cell for row in sheet.iter_rows() for cell in row
    }


def _formula_cells_without_value(data: bytes) -> dict[str, int]:
    """Count formula cells per sheet whose cached value is missing (never computed by Excel)."""
    counts: dict[str, int] = {}
    try:
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # pragma: no cover - the main load already succeeded
        return counts
    try:
        for name in formulas.sheetnames:
            missing = 0
            rows_f = formulas[name].iter_rows(values_only=True)
            rows_v = values[name].iter_rows(values_only=True)
            for row_f, row_v in zip(rows_f, rows_v, strict=False):
                for cell_f, cell_v in zip(row_f, row_v, strict=False):
                    if isinstance(cell_f, str) and cell_f.startswith("=") and cell_v is None:
                        missing += 1
            if missing:
                counts[name] = missing
    finally:
        formulas.close()
        values.close()
    return counts


def extract_xlsx(data: bytes, draft: Draft) -> None:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        properties = workbook.properties
        draft.title_hint = one_line(properties.title or "") or None
        created = properties.created
        draft.authored = created.date() if isinstance(created, datetime) else None
        draft.lang = lang_from_tag(getattr(properties, "language", None))
        sheets = list(workbook.worksheets)
        labels = [
            safe_label(sheet.title, fallback=f"sheet {number}", max_len=60)
            for number, sheet in enumerate(sheets, 1)
        ]
        slugs = assign_unique_slugs([sheet.title for sheet in sheets], max_len=40, fallback="sheet")
        for sheet, label, slug in zip(sheets, labels, slugs, strict=True):
            _extract_sheet(sheet, label, slug, draft)
    finally:
        workbook.close()
    by_name = dict(zip((sheet.title for sheet in sheets), labels, strict=True))
    for name, missing in sorted(_formula_cells_without_value(data).items()):
        sheet_label = by_name.get(name) or safe_label(name, fallback="sheet", max_len=60)
        draft.warn(
            "xlsx-formula-no-value",
            f'sheet "{sheet_label}": {missing} formula cell(s) have no cached value (left empty)',
        )


def _extract_sheet(sheet: Any, label: str, slug: str, draft: Draft) -> None:
    raw = _sheet_cells(sheet)
    values: dict[tuple[int, int], str] = {}
    for (row, column), cell in raw.items():
        if column > MAX_TABLE_COLUMNS or row > MAX_TABLE_ROWS + 1:
            continue
        text = format_cell(cell.value, getattr(cell, "number_format", None))
        if text:
            values[(row, column)] = text
    for merged in sheet.merged_cells.ranges:
        origin = values.get((merged.min_row, merged.min_col), "")
        if not origin:
            continue
        for row in range(merged.min_row, min(merged.max_row, MAX_TABLE_ROWS + 1) + 1):
            for column in range(merged.min_col, min(merged.max_col, MAX_TABLE_COLUMNS) + 1):
                values[(row, column)] = origin
    if not values:
        return
    last_row = max(row for row, _ in values)
    last_column = max(column for _, column in values)
    grid: dict[int, list[str]] = {}
    for (row, column), text in values.items():
        grid.setdefault(row, [""] * last_column)[column - 1] = text
    header = grid.get(1, [""] * last_column)
    rows = [(row, grid[row]) for row in sorted(grid) if row > 1]

    default_slug = slugify_ascii(label, 40) or "sheet"
    sheet_slug = slug if slug != default_slug else None

    def locator(start: int, end: int) -> XlsxRowsLocator:
        return XlsxRowsLocator(sheet=label, sheet_slug=sheet_slug, row_start=start, row_end=end)

    draft.block("heading", label, locator(1, 1), "deterministic-table", level=2)
    emit_table(
        draft,
        header=header,
        rows=rows,
        locator=locator(1, last_row),
        provenance="deterministic-table",
        row_locator=lambda row: locator(row, row),
        what=f'sheet "{label}"',
    )
