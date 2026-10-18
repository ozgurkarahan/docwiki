"""DOCX via python-docx: body paragraphs, tables and inline pictures (local adapter).

Locators follow python-docx indices: ``docx-paragraph`` N is ``Document.paragraphs[N - 1]``
(body paragraphs only, empty ones included; table cells, headers and footers excluded) and
``docx-table`` T is ``Document.tables[T - 1]`` with rows numbered from 1 (row 1 = header).
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from docwiki.contracts.enums import BlockKind
from docwiki.contracts.locators import DocxParagraphLocator, DocxTableLocator, Locator
from docwiki.extract._draft import Draft
from docwiki.extract._images import embedded_png
from docwiki.extract._tables import emit_table
from docwiki.extract._text import BULLET_RE, clean_text, lang_from_tag, one_line

__all__ = ["extract_docx"]

_HEADING_STYLE_RE = re.compile(r"^heading\s*([1-9])$", re.IGNORECASE)
_P = qn("w:p")
_TBL = qn("w:tbl")
_BLIP = qn("a:blip")
_EMBED = qn("r:embed")
_OUTLINE = qn("w:outlineLvl")
_NUMPR = qn("w:numPr")
_VAL = qn("w:val")


def _style_name(paragraph: Paragraph) -> str:
    try:
        style = paragraph.style
    except (KeyError, ValueError, AttributeError):
        return ""
    return (style.name or "") if style is not None else ""


def _kind_level(paragraph: Paragraph, text: str) -> tuple[BlockKind, int | None]:
    name = _style_name(paragraph).strip()
    match = _HEADING_STYLE_RE.match(name)
    if match:
        return "heading", int(match[1])
    lowered = name.lower()
    if lowered == "title":
        return "heading", 1
    if lowered == "subtitle":
        return "heading", 2
    ppr = paragraph._p.pPr
    if ppr is not None:
        outline = ppr.find(_OUTLINE)
        if outline is not None:
            value = outline.get(_VAL, "")
            if value.isdigit() and int(value) < 9:
                return "heading", int(value) + 1
        if ppr.find(_NUMPR) is not None:
            return "list", None
    if lowered.startswith("list") or BULLET_RE.match(text):
        return "list", None
    return "paragraph", None


def _table_grid(table: DocxTable) -> list[list[str]]:
    grid: list[list[str]] = []
    for row in table.rows:
        try:
            cells = row.cells
        except (IndexError, ValueError, AttributeError):
            cells = []
        grid.append([clean_text(cell.text) for cell in cells])
    return grid


def extract_docx(data: bytes, draft: Draft) -> None:
    document = Document(io.BytesIO(data))
    body = document._body
    part = document.part
    properties = document.core_properties
    title = one_line(properties.title or "")
    created = properties.created
    draft.authored = created.date() if created is not None else None
    draft.lang = lang_from_tag(properties.language)

    paragraph_number = table_number = 0
    first_heading: str | None = None
    style_title: str | None = None
    for child in document.element.body.iterchildren():
        if child.tag == _P:
            paragraph_number += 1
            paragraph = Paragraph(child, body)
            text = clean_text(paragraph.text)
            locator = DocxParagraphLocator(paragraph=paragraph_number)
            if text:
                kind, level = _kind_level(paragraph, text)
                draft.block(kind, text, locator, "native-text", level=level)
                if kind == "heading":
                    first_heading = first_heading or one_line(text)
                    if _style_name(paragraph).lower() == "title":
                        style_title = style_title or one_line(text)
            for blip in child.iter(_BLIP):
                rel_id = blip.get(_EMBED)
                related = part.related_parts.get(rel_id) if rel_id else None
                blob = getattr(related, "blob", None)
                if not blob:
                    continue
                png = embedded_png(blob)
                if png is None:
                    draft.warn(
                        "figure-skipped",
                        f"¶ {paragraph_number}: picture format not supported or too small",
                    )
                    continue
                draft.figure(locator, png)
        elif child.tag == _TBL:
            table_number += 1
            grid = _table_grid(DocxTable(child, body))
            if not grid:
                continue
            emit_table(
                draft,
                header=grid[0],
                rows=[(index, cells) for index, cells in enumerate(grid[1:], 2)],
                locator=DocxTableLocator(table=table_number, row_start=1, row_end=len(grid)),
                provenance="deterministic-table",
                row_locator=_row_locator(table_number),
                what=f"table {table_number}",
            )
    draft.title_hint = style_title or title or first_heading


def _row_locator(table: int) -> Callable[[int], Locator]:
    def locate(row: int) -> Locator:
        return DocxTableLocator(table=table, row_start=row, row_end=row)

    return locate
