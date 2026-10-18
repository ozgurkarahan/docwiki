from __future__ import annotations

import importlib.util
import io
import zipfile
from pathlib import Path

import pytest
from leakguard_core.extract import extract_bytes
from leakguard_core.model import Extraction
from leakguard_fixtures import (
    FAKE_EMAIL,
    FAKE_GUID,
    FAKE_INTERNAL_URL,
    FAKE_TERM,
    make_docx,
    make_eml,
    make_image,
    make_pdf,
    make_pptx,
    make_sqlite,
    make_xlsx,
    replace_zip_member,
)


def texts(extraction: Extraction) -> str:
    return "\n".join(chunk.text for chunk in extraction.chunks)


def meta(extraction: Extraction) -> dict[str, str]:
    return {field.name: field.value for field in extraction.metadata}


def test_docx_text_tables_metadata_and_company() -> None:
    data = make_docx(
        ["Gearbox oil change for " + FAKE_TERM, "Contact " + FAKE_EMAIL],
        author="Jane Quillon",
        last_modified_by="jquillon",
        table=[["Supplier", "Zorblax"]],
        company="Zorblax Holdings",
    )
    extraction = extract_bytes(data, "sop.docx")
    assert extraction.kind == "office"
    body = texts(extraction)
    assert FAKE_TERM in body and FAKE_EMAIL in body
    fields = meta(extraction)
    assert fields["creator"] == "Jane Quillon"
    assert fields["lastModifiedBy"] == "jquillon"
    assert fields["Company"] == "Zorblax Holdings"
    cells = [c for c in extraction.chunks if c.text == "Zorblax"]
    assert cells and not cells[0].sentence_start  # table cells are not sentence starts
    assert extraction.exif == []


def test_docx_default_thumbnail_carries_exif() -> None:
    extraction = extract_bytes(make_docx(["x"], keep_thumbnail=True), "x.docx")
    assert extraction.exif == ["docProps/thumbnail.jpeg > EXIF block"]


def test_docx_hidden_parts_are_extracted() -> None:
    data = make_docx(["Visible text"])
    comments = (
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:comment w:id="0" w:author="Quentin Marsh" w:initials="QM"><w:p><w:r><w:t>'
        "Ask " + FAKE_EMAIL + "</w:t></w:r></w:p></w:comment></w:comments>"
    )
    data = replace_zip_member(data, "word/comments.xml", comments.encode("utf-8"))
    extraction = extract_bytes(data, "x.docx")
    body = texts(extraction)
    assert "Quentin Marsh" in body and FAKE_EMAIL in body
    author = next(c for c in extraction.chunks if c.text == "Quentin Marsh")
    assert author.kind == "metadata" and not author.sentence_start


def test_xlsx_cells_formulas_and_metadata() -> None:
    data = make_xlsx(
        [["Part", "Supplier"], ["Gearbox oil", "Zorblax"], ["Filter", FAKE_GUID]],
        creator="Jane Quillon",
        formula="=HYPERLINK(\"" + FAKE_INTERNAL_URL + "\")",
    )
    extraction = extract_bytes(data, "stock.xlsx")
    locations = {chunk.text: chunk.location for chunk in extraction.chunks}
    assert locations["Zorblax"] == "Plan!B2"
    assert locations[FAKE_GUID] == "Plan!B3"
    formula = next(c for c in extraction.chunks if c.kind == "formula")
    assert FAKE_INTERNAL_URL in formula.text and not formula.closed_world
    assert meta(extraction)["creator"] == "Jane Quillon"


def test_pptx_text_and_default_last_modified_by() -> None:
    extraction = extract_bytes(make_pptx("Site Rules", "Escort by Zorblax staff"), "deck.pptx")
    body = texts(extraction)
    assert "Site Rules" in body and "Escort by Zorblax staff" in body
    assert meta(extraction)["creator"] == "DocWiki sample generator"
    library_default = extract_bytes(make_pptx("T", "B", last_modified_by=None), "deck.pptx")
    assert meta(library_default).get("lastModifiedBy")  # python-pptx ships a person's name


def test_pptx_table_style_reference_is_not_scanned() -> None:
    import pptx
    from leakguard_core.rules import GUID_RE
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Asset"
    buffer = io.BytesIO()
    presentation.save(buffer)
    raw = buffer.getvalue()
    assert b"tableStyleId" in zipfile.ZipFile(buffer).read("ppt/slides/slide1.xml")
    extraction = extract_bytes(raw, "deck.pptx")
    assert "Asset" in texts(extraction)
    assert not [chunk.text for chunk in extraction.chunks if GUID_RE.search(chunk.text)]


def test_pdf_text_metadata_and_links() -> None:
    data = make_pdf(
        ["Lab report for " + FAKE_TERM, "Sample id " + FAKE_GUID],
        {"Author": "jquillon", "Creator": "Quillon Writer 2", "Producer": "fpdf2"},
        uri=FAKE_INTERNAL_URL,
    )
    extraction = extract_bytes(data, "lab.pdf")
    assert extraction.kind == "pdf", extraction.problems
    body = texts(extraction)
    assert FAKE_TERM in body and FAKE_GUID in body and FAKE_INTERNAL_URL in body
    assert meta(extraction) == {
        "Author": "jquillon",
        "Creator": "Quillon Writer 2",
        "Producer": "fpdf2",
    }
    page_chunk = next(c for c in extraction.chunks if FAKE_TERM in c.text)
    assert page_chunk.location.startswith("page 1 line")


def test_png_text_chunks_and_exif() -> None:
    clean = extract_bytes(make_image("PNG", text={"Author": "Jane Quillon"}), "a.png")
    assert clean.exif == []
    assert meta(clean)["Author"] == "Jane Quillon"
    assert "image pixels are not OCR-scanned" in clean.notices
    with_exif = extract_bytes(make_image("PNG", exif=True), "b.png")
    assert with_exif.exif


def test_jpeg_exif_detection() -> None:
    assert extract_bytes(make_image("JPEG"), "a.jpg").exif == []
    assert extract_bytes(make_image("JPEG", exif=True), "b.jpg").exif == ["EXIF block"]


def test_sqlite_text_columns(tmp_path: Path) -> None:
    database = make_sqlite(tmp_path / "search.db", [("Gearbox", "Owner " + FAKE_EMAIL)])
    extraction = extract_bytes(database.read_bytes(), "search.db")
    assert extraction.kind == "sqlite", extraction.problems
    chunk = next(c for c in extraction.chunks if FAKE_EMAIL in c.text)
    assert chunk.location == "table notes row 1 column body"


@pytest.mark.skipif(importlib.util.find_spec("duckdb") is None, reason="duckdb not installed")
def test_parquet_rows(tmp_path: Path) -> None:
    import duckdb

    target = tmp_path / "kpi.parquet"
    connection = duckdb.connect()
    connection.execute(
        f"COPY (SELECT 'Turbine 7' AS asset, '{FAKE_TERM}' AS vendor) TO '{target.as_posix()}' "
        "(FORMAT parquet)"
    )
    connection.close()
    extraction = extract_bytes(target.read_bytes(), "kpi.parquet", path_on_disk=target)
    assert FAKE_TERM in texts(extraction)
    detached = extract_bytes(target.read_bytes(), "kpi.parquet")
    assert detached.chunks == [] and detached.notices[0].startswith("unscanned")


def test_eml_headers_body_and_attachment() -> None:
    data = make_eml(
        "Please ask " + FAKE_EMAIL + " about the bulletin.",
        attachment=("ids.txt", ("subscription " + FAKE_GUID).encode()),
    )
    extraction = extract_bytes(data, "mail.eml")
    body = texts(extraction)
    assert "Subject: Bulletin applicability" in body
    assert FAKE_EMAIL in body and FAKE_GUID in body
    attached = next(c for c in extraction.chunks if FAKE_GUID in c.text)
    assert attached.location.startswith("attachment ids.txt")


def test_unicode_escapes_are_decoded_for_every_rule() -> None:
    source = '{"term": "Generatorb\\u00fcrsten", "vendor": "Z\\u00f6rblax"}\n'
    extraction = extract_bytes(source.encode("ascii"), "glossary.json")
    decoded = [c for c in extraction.chunks if "Generatorb\u00fcrsten" in c.text]
    assert len(decoded) == 1 and decoded[0].closed_world and decoded[0].identifiers
    escaped = [c for c in extraction.chunks if "\\u00fc" in c.text]
    assert len(escaped) == 1 and not escaped[0].closed_world and escaped[0].identifiers
    assert decoded[0].location == escaped[0].location == "line 1"
    surrogate = extract_bytes(b'"smile \\ud83d\\ude00"', "x.json")
    assert [c.text for c in surrogate.chunks] == ['"smile \\ud83d\\ude00"']


def test_unknown_binary_is_scanned_for_strings() -> None:
    data = b"\x00\x01\x02" + ("contact " + FAKE_EMAIL).encode() + b"\x00\xff" * 10
    extraction = extract_bytes(data, "blob.bin")
    assert extraction.unknown_binary
    assert FAKE_EMAIL in texts(extraction)
    assert all(not chunk.closed_world for chunk in extraction.chunks)


def test_python_files_expose_literals_to_the_noun_check() -> None:
    source = 'import os\n\nclass ZorbHelper:\n    """Made for Quentin Marsh."""\n# Zorblax note\n'
    extraction = extract_bytes(source.encode(), "gen.py")
    prose = [c.text for c in extraction.chunks if c.closed_world]
    assert any("Quentin Marsh" in text for text in prose)
    assert any("Zorblax note" in text for text in prose)
    assert not any("ZorbHelper" in text for text in prose)


def test_markdown_code_fences_skip_the_noun_check() -> None:
    text = "# Title\n\n```python\nZorbClient()\n```\nBody text.\n"
    extraction = extract_bytes(text.encode(), "page.md")
    fenced = [c for c in extraction.chunks if "ZorbClient" in c.text]
    assert fenced and not fenced[0].closed_world
    heading = extraction.chunks[0]
    assert heading.text == "# Title" and not heading.sentence_start


def test_corrupt_office_file_reports_a_problem(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("LEAKGUARD_DEBUG", raising=False)
    extraction = extract_bytes(b"PK\x03\x04broken", "broken.docx")
    assert extraction.problems
