"""Fixture builders for the leak guard tests.

Sensitive-looking values are assembled at run time so that the test sources themselves never
contain a literal identifier, secret or path (the repository is scanned by the leak guard too).
"""

from __future__ import annotations

import io
import subprocess
import zipfile
from pathlib import Path

FAKE_TERM = "ZORBLAX-CORP"
FAKE_GUID = "-".join(["0f1e2d3c", "4b5a", "6978", "8a9b", "0c1d2e3f4a5b"])
FAKE_EMAIL = "jane.doe" + "@" + "zorblax-corp.com"
FAKE_AUTHOR_EMAIL = "dev" + "@" + "zorblax-corp.com"
FAKE_WIN_HOME = "C:" + "\\" + "Users" + "\\" + "zorblax" + "\\" + "notes.txt"
FAKE_UNIX_HOME = "/" + "home" + "/zorblax/project"
FAKE_MAC_HOME = "/" + "Users" + "/zorblax/project"
FAKE_GITHUB_TOKEN = "gh" + "p_" + "A1b2C3d4E5f6G7h8I9j0K1l2M3n4O5p6Q7r8"
FAKE_ACCOUNT_KEY = "Account" + "Key=" + "Abcdefghijklmnopqrstuvwxyz0123456789ABCD=="
FAKE_SHARED_KEY = "SharedAccess" + "Key=" + "Zyxwvutsrqponmlkjihgfedcba9876543210ZYXW="
FAKE_SAS = (
    "?sv=2024-01-01&sp=r&" + "sig" + "=" + "Q2FzZVNlbnNpdGl2ZVNpZ25hdHVyZVZhbHVl%2BAbC12%3D"
)
FAKE_JWT = (
    "eyJ" + "hbGciOiJIUzI1NiJ9" + "." + "eyJ" + "zdWIiOiJ6b3JibGF4In0" + "." + "c2lnbmF0dXJlMTIz"
)
FAKE_PRIVATE_KEY = "-----BEGIN " + "RSA PRIVATE KEY-----"
FAKE_API_KEY_ASSIGNMENT = "api" + "_key = '" + "Zx9Qv7Lm3Np5Rt" + "'"
FAKE_ENV_SECRET = "DOCWIKI_CLIENT" + "_SECRET=" + "Qw8Er7Ty6Ui5Op4"
FAKE_URL_CREDENTIALS = "https:" + "//builder:" + "Pa55wordValue" + "@git.zorblax-corp.com/repo"
FAKE_INTERNAL_URL = "https:" + "//wiki." + "corp" + ".zorblax-corp.com/page"
FAKE_SINGLE_LABEL_URL = "http:" + "//build01/job"
FAKE_BARE_INTERNAL = "planning01." + "lan"
FAKE_MIRROR_URL = "https:" + "//pkgs.dev.azure.com/zorblax/_packaging/feed/pypi/simple/"
FAKE_INDEX_LINE = "--index" + "-url " + "https:" + "//packages.zorblax-corp.com/repo"
PUBLIC_REGISTRY_URL = "https:" + "//pypi.org/simple/"


def git(repo: Path, *args: str) -> str:
    """Run git in a test repository without touching any user or global configuration."""
    command = [
        "git", "-C", str(repo),
        "-c", "user.name=DocWiki Test", "-c", "user.email=test@example.com",
        "-c", "commit.gpgsign=false", "-c", "tag.gpgsign=false", "-c", "core.autocrlf=false",
        "-c", "init.defaultBranch=main", "-c", f"core.hooksPath={repo / '.no-hooks'}",
        *args,
    ]
    completed = subprocess.run(command, capture_output=True, text=True, check=True)
    return completed.stdout


def commit_all(repo: Path, message: str, *, author_email: str | None = None) -> str:
    git(repo, "add", "-A")
    extra = ["-c", f"user.email={author_email}"] if author_email else []
    command = [*extra, "commit", "-q", "--no-verify", "-m", message]
    git(repo, *command)
    return git(repo, "rev-parse", "HEAD").strip()


def write(repo: Path, relative: str, content: str | bytes) -> Path:
    path = repo / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8", newline="\n")
    return path


def replace_zip_member(data: bytes, name: str, content: bytes | None) -> bytes:
    """Return a copy of a ZIP (Office) file with one member replaced (or removed if None)."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename == name:
                continue
            target.writestr(info, source.read(info))
        if content is not None:
            target.writestr(name, content)
    return output.getvalue()


def _strip_thumbnail(data: bytes) -> bytes:
    return replace_zip_member(data, "docProps/thumbnail.jpeg", None)


def make_docx(
    paragraphs: list[str],
    *,
    author: str = "DocWiki sample generator",
    last_modified_by: str = "",
    table: list[list[str]] | None = None,
    company: str | None = None,
    keep_thumbnail: bool = False,
) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                grid.cell(row_index, col_index).text = value
    properties = document.core_properties
    properties.author = author
    properties.last_modified_by = last_modified_by
    properties.comments = ""
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    if company is not None:
        app = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
            'extended-properties"><Application>DocWiki</Application>'
            f"<Company>{company}</Company><Manager></Manager></Properties>"
        )
        data = replace_zip_member(data, "docProps/app.xml", app.encode("utf-8"))
    return data if keep_thumbnail else _strip_thumbnail(data)


def make_xlsx(
    rows: list[list[object]],
    *,
    creator: str = "DocWiki sample generator",
    sheet: str = "Planning",
    formula: str | None = None,
) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    if formula:
        worksheet["Z1"] = formula
    workbook.properties.creator = creator
    workbook.properties.lastModifiedBy = ""
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx(
    title: str,
    body: str,
    *,
    author: str = "DocWiki sample generator",
    last_modified_by: str | None = "",
    keep_thumbnail: bool = False,
) -> bytes:
    import pptx

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    presentation.core_properties.author = author
    if last_modified_by is not None:
        presentation.core_properties.last_modified_by = last_modified_by
    buffer = io.BytesIO()
    presentation.save(buffer)
    data = buffer.getvalue()
    return data if keep_thumbnail else _strip_thumbnail(data)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: list[str], info: dict[str, str], uri: str | None = None) -> bytes:
    """A minimal one-page PDF with a Helvetica text layer, an Info dictionary and a link."""
    operations = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
    operations.extend(f"({_pdf_escape(line)}) Tj T*" for line in lines)
    operations.append("ET")
    stream = "\n".join(operations).encode("latin-1")
    annotations = " /Annots [6 0 R]" if uri else ""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font "
            f"<< /F1 4 0 R >> >> /Contents 5 0 R{annotations} >>"
        ).encode("latin-1"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    if uri:
        objects.append(
            (
                "<< /Type /Annot /Subtype /Link /Rect [72 700 300 720] /Border [0 0 0] "
                f"/A << /S /URI /URI ({_pdf_escape(uri)}) >> >>"
            ).encode("latin-1")
        )
    entries = " ".join(f"/{key} ({_pdf_escape(value)})" for key, value in info.items())
    objects.append(f"<< {entries} >>".encode("latin-1"))
    info_number = len(objects)
    output = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        output += f"{offset:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info {info_number} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(output)


def make_image(
    fmt: str = "PNG", *, text: dict[str, str] | None = None, exif: bool = False
) -> bytes:
    from PIL import Image, PngImagePlugin

    image = Image.new("RGB", (4, 4), "white")
    options: dict[str, object] = {}
    if fmt == "PNG" and text:
        info = PngImagePlugin.PngInfo()
        for key, value in text.items():
            info.add_text(key, value)
        options["pnginfo"] = info
    if exif:
        data = Image.Exif()
        data[0x0131] = "ZorblaxCam firmware"  # Software tag
        options["exif"] = data.tobytes()
    buffer = io.BytesIO()
    image.save(buffer, fmt, **options)
    return buffer.getvalue()


def make_sqlite(path: Path, rows: list[tuple[str, str]]) -> Path:
    import sqlite3

    connection = sqlite3.connect(path)
    with connection:
        connection.execute("CREATE TABLE notes (title TEXT, body TEXT)")
        connection.executemany("INSERT INTO notes VALUES (?, ?)", rows)
    connection.close()
    return path


def make_eml(body: str, *, attachment: tuple[str, bytes] | None = None, sender: str = "") -> bytes:
    from email.message import EmailMessage

    message = EmailMessage()
    message["From"] = sender or "Nadia Ferrandon <nadia.ferrandon@zava-materiaux.example>"
    message["To"] = "ines.roux@lamna-batiment.example"
    message["Subject"] = "Window delivery date"
    message.set_content(body)
    if attachment:
        name, content = attachment
        message.add_attachment(
            content, maintype="application", subtype="octet-stream", filename=name
        )
    return message.as_bytes()
