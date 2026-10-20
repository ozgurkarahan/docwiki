"""Local adapters: locators, labels, provenance, tables, figures, warnings, determinism."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from extract_fixtures import CORPUS_FILES, Corpus, make_text_pdf
from openpyxl import load_workbook

from docwiki.contracts import ExtractedDoc, locator_label, source_id_for
from docwiki.contracts.locators import (
    CsvRowsLocator,
    DocxParagraphLocator,
    DocxTableLocator,
    EmlPartLocator,
    ImageRegionLocator,
    PdfPageLocator,
    PptxSlideLocator,
    TextLinesLocator,
    VttCueLocator,
    XlsxRowsLocator,
)
from docwiki.engine_types import ExtractionError, ExtractSettings

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def keys(doc: ExtractedDoc) -> list[str]:
    return [block.locator_key for block in doc.blocks]


async def test_pdf_paragraphs_headings_figures_and_page_images(corpus: Corpus) -> None:
    result = await corpus.extract("docs/devis-indice-a.pdf")
    doc = result.doc
    assert doc.source_id == source_id_for("docs/devis-indice-a.pdf")
    assert doc.source_hash.startswith("sha256:")
    assert doc.media_type == "application/pdf"
    assert doc.metadata.page_count == 2
    assert doc.title_hint == "Devis DEV-2026-014 indice A"
    assert str(doc.metadata.authored) == "2026-01-15"
    assert doc.extraction.adapter == "local"
    assert doc.extraction.warnings == [] == result.warnings

    texts = [block.text for block in doc.blocks]
    assert texts[0] == "Devis de l'extension"
    assert doc.blocks[0].kind == "heading" and doc.blocks[0].level == 1
    assert any("FICTIONAL - NOT FOR OPERATIONAL USE" in text for text in texts)
    assert "Nettoyer la dalle avant la pose." in texts
    lists = [block for block in doc.blocks if block.kind == "list"]
    assert [block.text for block in lists] == [
        "1. Fermer l'ouverture provisoire.",
        "2. Contrôler la livraison.",
    ]
    for block in doc.blocks:
        assert isinstance(block.locator, PdfPageLocator)
        assert block.provenance == "native-text"
        assert block.locator.bbox is not None
        x0, y0, x1, y1 = block.locator.bbox
        assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1
    assert [b.id for b in doc.blocks] == [f"b{i:04d}" for i in range(1, len(doc.blocks) + 1)]
    assert {block.locator_key for block in doc.blocks} == {"p0001", "p0002"}
    assert locator_label(doc.blocks[0].locator) == "p. 1"

    stamp = next(b for b in doc.blocks if b.text.startswith("FICTIONAL"))
    paragraph = next(b for b in doc.blocks if b.text.startswith("Toutes les baies"))
    assert stamp.locator.bbox[1] < paragraph.locator.bbox[1]  # top-left origin, reading order

    assert [figure.id for figure in doc.figures] == ["f01"]
    figure = doc.figures[0]
    assert figure.image == "figures/f01.png"
    assert figure.locator == PdfPageLocator(page=1, bbox=figure.bbox)
    assert figure.bbox == pytest.approx([20 / 210, 150 / 297, 90 / 210, 193.75 / 297], abs=0.01)
    assert figure.description is None and figure.description_provenance is None
    assert result.figures["f01"].startswith(PNG_SIGNATURE)
    assert sorted(result.page_images) == [1, 2]
    assert all(png.startswith(PNG_SIGNATURE) for png in result.page_images.values())


async def test_scanned_pdf_needs_ocr(corpus: Corpus) -> None:
    result = await corpus.extract("docs/bon-de-livraison-scan.pdf")
    doc = result.doc
    assert len(doc.blocks) == 1
    block = doc.blocks[0]
    assert (block.kind, block.text, block.provenance) == ("ocr-line", "", "ocr")
    assert block.locator == PdfPageLocator(page=1)
    assert any(warning.startswith("needs-ocr:") for warning in result.warnings)
    assert doc.figures == []  # a full-page scan is not a figure
    assert 1 in result.page_images


async def test_pdf_max_pages(corpus: Corpus, tmp_path: Path) -> None:
    path = tmp_path / "long.pdf"
    make_text_pdf(path, pages=3)
    result = await corpus.extract("docs/long.pdf", ExtractSettings(max_pages=2), path=path)
    assert result.doc.metadata.page_count == 3
    assert sorted(result.page_images) == [1, 2]
    assert {block.locator.page for block in result.doc.blocks} == {1, 2}
    assert any(w.startswith("max-pages: extracted the first 2 of 3 pages") for w in result.warnings)


async def test_docx_paragraph_indices_follow_python_docx(corpus: Corpus) -> None:
    result = await corpus.extract("docs/avenant-n1.docx")
    doc = result.doc
    paragraphs = Document(str(corpus.path("docs/avenant-n1.docx"))).paragraphs
    text_blocks = [b for b in doc.blocks if isinstance(b.locator, DocxParagraphLocator)]
    assert [b.locator.paragraph for b in text_blocks] == [1, 2, 4, 5, 7, 9]
    for block in text_blocks:
        assert paragraphs[block.locator.paragraph - 1].text == block.text
    assert paragraphs[2].text == "" and paragraphs[5].text == ""  # empty ones keep their index
    title, scope = text_blocks[0], text_blocks[2]
    assert (title.kind, title.level) == ("heading", 1)
    assert (scope.kind, scope.level, scope.text) == ("heading", 1, "Objet")
    assert text_blocks[4].kind == "list"
    assert keys(doc)[4:6] == ["tbl-01-r0002-r0002", "tbl-01-r0003-r0003"]
    assert doc.blocks[4].text == "Repère: FT1; Désignation: Fenêtre de toit; Montant HT: 1160,00 €"
    assert doc.blocks[4].provenance == "deterministic-table"
    assert locator_label(doc.blocks[4].locator) == "table 1, row 2"

    (table,) = doc.tables
    assert table.id == "t01" and table.artifact == "tables/t01.parquet"
    assert table.locator == DocxTableLocator(table=1, row_start=1, row_end=3)
    assert table.columns == ["Repère", "Désignation", "Montant HT"]
    assert [row.row_id for row in table.rows] == ["r0002", "r0003"]
    assert table.provenance == "deterministic-table"

    (figure,) = doc.figures
    assert figure.locator == DocxParagraphLocator(paragraph=8)
    assert paragraphs[7].text == ""  # the picture paragraph
    assert result.figures["f01"].startswith(PNG_SIGNATURE)


async def test_xlsx_values_rows_slugs_and_merges(corpus: Corpus) -> None:
    result = await corpus.extract("plans/planning-chantier.xlsx")
    doc = result.doc
    planning, notes = doc.tables
    assert planning.locator == XlsxRowsLocator(sheet="Planning", row_start=1, row_end=6)
    assert planning.columns == ["Tâche", "Désignation", "Durée", "Début", "Avancement", "Double"]
    rows = {row.row_id: row.cells for row in planning.rows}
    assert list(rows) == ["r0002", "r0003", "r0004", "r0006"]  # Excel numbers, empty row 5 skipped
    assert rows["r0002"] == ["T05", "Élévation des murs", "6", "2026-03-17", "20%", "12"]
    assert rows["r0003"][2:4] == ["2.5", "2026-03-30 13:30:00"]
    assert rows["r0006"][:3] == ["Remarque fusionnée", "Remarque fusionnée", "1"]  # A6:B6 merged
    assert all("=" not in cell for cells in rows.values() for cell in cells)  # values, not formulas
    assert planning.provenance == notes.provenance == "deterministic-table"
    assert notes.locator.sheet == "Notes & R\u00e9serves"
    assert {block.provenance for block in doc.blocks} == {"deterministic-table"}
    assert keys(doc) == [
        "sheet-planning-r0001-r0001",
        "sheet-planning-r0002-r0002",
        "sheet-planning-r0003-r0003",
        "sheet-planning-r0004-r0004",
        "sheet-planning-r0006-r0006",
        "sheet-notes-reserves-r0001-r0001",
        "sheet-notes-reserves-r0002-r0002",
    ]
    assert locator_label(doc.blocks[1].locator) == 'sheet "Planning", row 2'
    assert doc.blocks[0].kind == "heading" and doc.blocks[0].text == "Planning"
    # the workbook really stores a formula with a cached value
    sheet = load_workbook(corpus.path("plans/planning-chantier.xlsx"))["Planning"]
    assert sheet["F2"].value == "=C2*2"


async def test_xlsx_colliding_sheet_slugs_and_formula_without_value(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from docwiki.extract import extract_source

    workbook = Workbook()
    first = workbook.active
    first.title = "Planning"
    first.append(["A", "B"])
    first.append([1, "=A2*2"])
    second = workbook.create_sheet("planning!")
    second.append(["A"])
    second.append(["x"])
    third = workbook.create_sheet("Planning{2026}")
    third.append(["A"])
    third.append(["y"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    result = await extract_source(path, rel_path="book.xlsx", settings=ExtractSettings())
    locators = [table.locator for table in result.doc.tables]
    # slugs come from slugify_ascii(raw sheet name); labels drop Markdown-significant characters
    assert [loc.sheet for loc in locators] == ["Planning", "planning!", "Planning2026"]
    assert [loc.sheet_slug for loc in locators] == [None, "planning-2", "planning-2026"]
    assert [b.locator_key for b in result.doc.blocks if b.kind == "table-ref"] == [
        "sheet-planning-r0002-r0002",
        "sheet-planning-2-r0002-r0002",
        "sheet-planning-2026-r0002-r0002",
    ]
    assert any(w.startswith('xlsx-formula-no-value: sheet "Planning": 1') for w in result.warnings)
    assert result.doc.tables[0].rows[0].cells == ["1", ""]


def test_extractor_versions() -> None:
    from docwiki.extract import ADAPTER_VERSION, CU_API_VERSION, extractor_versions

    versions = extractor_versions()
    assert versions["local"] == ADAPTER_VERSION
    assert versions["content-understanding"].endswith(CU_API_VERSION)
    assert {"pypdfium2", "python-docx", "openpyxl", "python-pptx", "pillow"} <= set(versions)


async def test_csv_bom_dialect_and_record_numbers(corpus: Corpus) -> None:
    doc = (await corpus.extract("plans/commandes.csv")).doc
    (table,) = doc.tables
    assert table.columns == ["Commande", "Désignation", "Montant HT"]  # BOM stripped, ';' sniffed
    assert table.locator == CsvRowsLocator(row_start=1, row_end=4)
    assert [(row.row_id, row.cells) for row in table.rows] == [
        ("r0002", ["BC-2026-009", "Baie coulissante F1\navec quincaillerie", "3320,00"]),
        ("r0004", ["BC-2026-010", "Fenêtre F2 PVC", "1020,00"]),
    ]
    assert keys(doc) == ["r0002-r0002", "r0004-r0004"]
    assert locator_label(doc.blocks[1].locator) == "row 4"
    assert doc.blocks[0].text == (
        "Commande: BC-2026-009; Désignation: Baie coulissante F1 avec quincaillerie; "
        "Montant HT: 3320,00"
    )


async def test_pptx_slides_notes_tables_pictures(corpus: Corpus) -> None:
    result = await corpus.extract("slides/support-reunion-chantier.pptx")
    doc = result.doc
    assert doc.metadata.page_count == 3
    assert doc.title_hint == "Réunion de chantier n°3"
    assert all(isinstance(block.locator, PptxSlideLocator) for block in doc.blocks)
    assert keys(doc) == ["slide-01"] * 2 + ["slide-02"] * 3 + ["slide-03"] * 3
    kinds = [(block.kind, block.text) for block in doc.blocks]
    assert kinds[2:5] == [
        ("heading", "Décisions"),
        ("slide", "Fermer les ouvertures\nAvancer la cuisine"),
        ("note", "Rappeler le report de livraison des menuiseries."),
    ]
    assert doc.blocks[4].locator == PptxSlideLocator(slide=2, shape="notes")
    assert kinds[6] == ("table-ref", "Tâche: T07; Date: 2026-04-20")
    (table,) = doc.tables
    assert table.provenance == "native-text"
    assert [row.row_id for row in table.rows] == ["r0002", "r0003"]
    (figure,) = doc.figures
    assert isinstance(figure.locator, PptxSlideLocator) and figure.locator.slide == 3
    assert locator_label(figure.locator) == "slide 3"


async def test_eml_headers_body_lines_and_attachments(corpus: Corpus) -> None:
    doc = (await corpus.extract("mail/report-de-livraison.eml")).doc
    assert doc.title_hint == "Report de livraison des menuiseries"
    assert str(doc.metadata.authored) == "2026-03-18"
    header = doc.blocks[0]
    assert header.kind == "email-header"
    assert header.locator == EmlPartLocator(part="headers")
    assert header.text.splitlines()[0] == (
        "From: Erwan Gouzelin <erwan.gouzelin@vanarsdel-menuiseries.example>"
    )
    assert "Subject: Report de livraison des menuiseries" in header.text
    body = [b for b in doc.blocks if b.kind == "email-body"]
    assert [(b.locator.line_start, b.locator.line_end) for b in body] == [(1, 1), (3, 4), (6, 7)]
    assert body[1].text == (
        "La livraison de la commande BC-2026-009 est reportée au lundi 20 avril 2026.\n"
        "Pénurie de profilés aluminium."
    )
    assert locator_label(body[1].locator) == "message body, lines 3–4"
    attachment = doc.blocks[-1]
    assert attachment.locator == EmlPartLocator(part="attachment:bon-de-commande.pdf")
    assert attachment.locator_key == "msg-att-bon-de-commande-pdf"
    assert attachment.text.startswith("Attachment: bon-de-commande.pdf (application/pdf")
    assert "alert" not in doc.text()


async def test_eml_html_only_body_is_converted_safely(corpus: Corpus) -> None:
    result = await corpus.extract("mail/html-only.eml")
    body = [b.text for b in result.doc.blocks if b.kind == "email-body"]
    assert body == [
        "Bonjour,",
        "La livraison de la commande BC-2026-009 est reportée au 20/04/2026.",
    ]
    assert "alert" not in result.doc.text() and "color" not in result.doc.text()
    assert any(w.startswith("eml-html-body") for w in result.warnings)


async def test_vtt_cues(corpus: Corpus) -> None:
    doc = (await corpus.extract("meetings/reunion-de-chantier.vtt")).doc
    first, second = doc.blocks
    assert first.locator == VttCueLocator(cue_start="00:00:01", cue_end="00:00:04")
    assert first.text == "Chef d'équipe: Livraison des palettes contrôlée sur le chantier."
    assert second.locator == VttCueLocator(cue_start="00:03:15", cue_end="00:03:42")
    assert second.text == "Fermer les ouvertures & noter la date\navant le 07/04."
    assert (first.locator_key, second.locator_key) == ("cue-00-00-01", "cue-00-03-15")
    assert locator_label(second.locator) == "03:15–03:42"
    assert {b.kind for b in doc.blocks} == {"cue"}


@pytest.mark.parametrize(
    ("rel", "first_kind"),
    [
        ("notes/consignes-de-pose.md", "heading"),
        ("notes/consignes-de-pose.txt", "paragraph"),
    ],
)
async def test_text_line_ranges(corpus: Corpus, rel: str, first_kind: str) -> None:
    doc = (await corpus.extract(rel)).doc
    assert [b.locator for b in doc.blocks] == [
        TextLinesLocator(line_start=1, line_end=1),
        TextLinesLocator(line_start=3, line_end=4),
        TextLinesLocator(line_start=7, line_end=7),
    ]
    assert doc.blocks[0].kind == first_kind
    assert doc.blocks[1].text == "La baie F1 mesure 3,00 x 2,15 m.\nNoter la mesure."
    assert doc.blocks[2].kind == "list"
    assert locator_label(doc.blocks[1].locator) == "lines 3–4"


async def test_image_figure_and_ocr_placeholder(corpus: Corpus) -> None:
    result = await corpus.extract("photos/photo-palettes.png")
    doc = result.doc
    region = ImageRegionLocator(region=1, bbox=[0.0, 0.0, 1.0, 1.0])
    (block,) = doc.blocks
    assert (block.kind, block.text, block.provenance, block.locator) == (
        "ocr-line",
        "",
        "ocr",
        region,
    )
    assert block.locator_key == "img-b01"
    (figure,) = doc.figures
    assert figure.locator == region
    assert result.page_images[1].startswith(PNG_SIGNATURE)
    assert any(w.startswith("needs-ocr") for w in result.warnings)


async def test_unsupported_and_damaged_files_raise(tmp_path: Path) -> None:
    from docwiki.extract import extract_source

    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ExtractionError) as unsupported:
        await extract_source(legacy, rel_path="old.doc", settings=ExtractSettings())
    assert unsupported.value.source_id == source_id_for("old.doc")
    assert "unsupported file type .doc" in unsupported.value.reason

    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"this is not a PDF")
    with pytest.raises(ExtractionError, match="not a readable PDF"):
        await extract_source(broken, rel_path="broken.pdf", settings=ExtractSettings())

    not_zip = tmp_path / "broken.docx"
    not_zip.write_bytes(b"plain text")
    with pytest.raises(ExtractionError, match="docx could not be parsed"):
        await extract_source(not_zip, rel_path="broken.docx", settings=ExtractSettings())

    with pytest.raises(ExtractionError, match="cannot read the file"):
        await extract_source(
            tmp_path / "missing.txt", rel_path="missing.txt", settings=ExtractSettings()
        )


async def test_invalid_settings_raise_value_error(corpus: Corpus) -> None:
    with pytest.raises(ValueError, match="max_pages"):
        await corpus.extract("notes/consignes-de-pose.txt", ExtractSettings(max_pages=0))
    with pytest.raises(ValueError, match="cu_endpoint"):
        await corpus.extract("docs/devis-indice-a.pdf", ExtractSettings(mode="azure"))


@pytest.mark.parametrize("rel", CORPUS_FILES)
async def test_extraction_is_deterministic(corpus: Corpus, rel: str) -> None:
    first = await corpus.extract(rel)
    second = await corpus.extract(rel)
    assert first.doc.to_json() == second.doc.to_json()
    assert first.figures == second.figures
    assert first.page_images == second.page_images
    assert ExtractedDoc.from_json(first.doc.to_json()) == first.doc


def test_docx_fixture_is_a_real_docx(corpus: Corpus) -> None:
    data = corpus.path("docs/avenant-n1.docx").read_bytes()
    assert len(Document(io.BytesIO(data)).tables) == 1
