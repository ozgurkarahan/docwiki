"""XLSX workbooks: planning and purchase tracking (indice A and the held-back indice B) and the
restricted internal price study.

Row 1 of every sheet is the table header (the extractor reads row 1 as the header and locators
count rows from it). Values only, no formulas: amounts are numbers shown with two decimals and
the euro sign, dates are date cells shown JJ/MM/AAAA. The stamp is in the print header and
footer of every sheet.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import STAMP, eur, fr_date, meta_for, percent
from .datasets import planning_tasks, price_study, purchases
from .ooxml import finalize_package
from .world import World

_HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
_AMOUNT_FORMAT = '0.00" €"'
_DATE_FORMAT = "DD/MM/YYYY"


def _value(value: object) -> object:
    return float(value) if isinstance(value, Decimal) else value


class WorkbookBuilder:
    def __init__(self, title: str) -> None:
        self.wb = Workbook()
        self.title = title
        self._first = True

    def sheet(
        self,
        name: str,
        header: Sequence[str],
        rows: Sequence[Sequence[object]],
        widths: Sequence[int],
        bold_rows: Sequence[int] = (),
    ) -> None:
        """Header in row 1, then the rows (``None`` = empty cell, ``[]`` = empty row)."""
        if self._first:
            ws = self.wb.active
            ws.title = name
            self._first = False
        else:
            ws = self.wb.create_sheet(name)
        ws.append(list(header))
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, row in enumerate(rows, start=2):
            for col, raw in enumerate(row, start=1):
                if raw is None:
                    continue
                cell = ws.cell(row=index, column=col, value=_value(raw))
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(raw, str))
                if isinstance(raw, Decimal):
                    cell.number_format = _AMOUNT_FORMAT
                elif isinstance(raw, date):
                    cell.number_format = _DATE_FORMAT
                if index - 1 in bold_rows:
                    cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.page_setup.orientation = "landscape"
        ws.oddHeader.left.text = self.title
        ws.oddHeader.right.text = STAMP
        ws.oddFooter.center.text = STAMP
        ws.oddFooter.right.text = "Page &P sur &N"

    def render(self, w: World, doc_id: str) -> bytes:
        doc = w.doc(doc_id)
        meta = meta_for(doc, w.org_name(str(doc["organisation"])))
        buf = io.BytesIO()
        self.wb.save(buf)
        return finalize_package(buf.getvalue(), meta)


# ------------------------------------------------------------------------------------------
# Planning and purchase tracking PLN-2026-003


def _planning(w: World, indice: str, doc_id: str) -> bytes:
    facts = w.facts["planning"]
    data = facts[f"indice_{indice}"]
    revision = indice.upper()
    reference = f"{facts['title']} {facts['number']} indice {revision}"
    b = WorkbookBuilder(f"{w.builder['name']} - {reference}")
    author = str(facts["author"])
    tasks = planning_tasks(w, indice)
    rows: list[Sequence[object]] = [
        [task.id, task.task, task.lot, task.by, task.start, task.end, percent(task.progress),
         task.note or None]
        for task in tasks
    ]
    rows.append([])
    rows += [
        ["Document", f"{reference} du {fr_date(str(data['date']))}"],
        ["Établi par", f"{w.name(author)}, {w.role(author).lower()}, "
                       f"{w.org_name(w.person(author)['organisation'])}"],
        ["Avancement au", date.fromisoformat(str(data["progress_date"]))],
    ]
    for key in ("supersedes_sentence", "reason_sentence"):
        if key in data:
            rows.append(["Observations", " ".join(str(data[key]).split())])
    b.sheet("Planning", facts["columns_planning"], rows, widths=[8, 40, 16, 24, 12, 12, 12, 48])
    orders, ordered, invoiced = purchases(w, indice)
    achats: list[Sequence[object]] = [
        [row.order, row.date, row.supplier, row.object, row.ordered_ht, row.invoiced_ht,
         row.status]
        for row in orders
    ]
    achats += [
        ["Total commandé (HT)", None, None, None, ordered, None, None],
        [str(facts["purchase_totals_label"]), None, None, None, None, invoiced, None],
    ]
    b.sheet("Achats", facts["columns_achats"], achats, widths=[26, 12, 24, 44, 14, 14, 52],
            bold_rows=(len(orders) + 1, len(orders) + 2))
    if b.wb.sheetnames != list(facts["sheets"]):
        raise ValueError(f"planning sheets {b.wb.sheetnames} differ from the bible")
    return b.render(w, doc_id)


def planning_a(w: World) -> bytes:
    return _planning(w, "a", "planning-a")


def planning_b(w: World) -> bytes:
    return _planning(w, "b", "planning-b")


# ------------------------------------------------------------------------------------------
# Restricted internal price study (basis: quote indice B)


def etude_prix(w: World) -> bytes:
    facts = w.facts["etude_de_prix"]
    doc = w.doc("etude-prix")
    lots, sums = price_study(w)
    b = WorkbookBuilder(str(doc["title"]))
    rows: list[Sequence[object]] = [
        [lot.lot, lot.title, lot.sale_ht, lot.materials_ht, lot.hours, lot.labour_ht,
         lot.subcontract_ht, lot.direct_cost_ht, lot.gross_margin_ht]
        for lot in lots
    ]
    rows.append(["Total", None, sums["sale_ht"], sums["materials_ht"], int(sums["hours"]),
                 sums["labour_ht"], sums["subcontract_ht"], sums["direct_cost_ht"],
                 sums["gross_margin_ht"]])
    rows.append([])
    hourly = f"{eur(sums['hourly_cost'])}/h"
    overheads = percent(facts["overheads_percent"])
    totals = facts["totals"]
    synthesis: list[tuple[str, object]] = [
        ("Base", f"{facts['basis']} du {fr_date(str(w.facts['devis']['indice_b']['date']))}"),
        ("Date de l'étude", date.fromisoformat(str(facts["date"]))),
        ("Coût horaire de main-d'œuvre", hourly),
        ("Matériaux", sums["materials_ht"]),
        (f"Main-d'œuvre ({int(sums['hours'])} h x {hourly})", sums["labour_ht"]),
        ("Sous-traitance", sums["subcontract_ht"]),
        ("Déboursé sec", sums["direct_cost_ht"]),
        (f"Frais généraux ({overheads} du déboursé sec)", sums["overheads_ht"]),
        ("Prix de revient", sums["cost_price_ht"]),
        ("Prix de vente", sums["sale_ht"]),
        ("Marge brute (prix de vente - déboursé sec)", sums["gross_margin_ht"]),
        ("Marge nette (prix de vente - prix de revient)", sums["net_margin_ht"]),
        ("Taux de marge nette", str(totals["net_margin_percent"])),
        ("Coefficient de vente (prix de vente / déboursé sec)", str(totals["sale_coefficient"])),
    ]
    rows.append([None, "Synthèse (montants HT)", "Valeur"])
    rows += [[None, label, value] for label, value in synthesis]
    b.sheet(
        str(facts["sheet"]),
        ["Lot", "Désignation", "Prix de vente HT", "Matériaux HT", "Heures", "Main-d'œuvre HT",
         "Sous-traitance HT", "Déboursé sec HT", "Marge brute HT"],
        rows,
        widths=[7, 48, 16, 14, 9, 16, 16, 16, 15],
        bold_rows=(len(lots) + 1, len(lots) + 3),
    )
    return b.render(w, "etude-prix")
