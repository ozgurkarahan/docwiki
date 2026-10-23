"""DOCX documents: the signed works contract, amendment n°1 and the site meeting minutes n°3."""

from __future__ import annotations

import io
from collections.abc import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .common import STAMP, eur, fr_date, meta_for, percent, quantity
from .datasets import check_amendment, check_contract
from .ooxml import finalize_package
from .world import World


def _sentence(text: str) -> str:
    return text[:1].upper() + text[1:]


def _flat(text: str) -> str:
    """Folded YAML text as one line."""
    return " ".join(str(text).split())


class DocxBuilder:
    """Small wrapper over python-docx that never creates empty body paragraphs."""

    def __init__(self, header: str, footer: str, legal: str = "") -> None:
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.0))
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        for lang in self.doc.styles.element.iter(qn("w:lang")):
            lang.set(qn("w:val"), "fr-FR")
        head = section.header.paragraphs[0]
        head.text = f"{header}  |  {STAMP}"
        head.runs[0].font.size = Pt(8)
        foot = section.footer.paragraphs[0]
        foot.text = f"{STAMP}  |  {footer}"
        foot.runs[0].font.size = Pt(8)
        if legal:
            # Same paragraph as the stamp: every header and footer paragraph carries the stamp.
            line = foot.add_run()
            line.add_break()
            line.add_text(legal)
            line.font.size = Pt(7)
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def title(self, text: str) -> None:
        self.doc.add_heading(text, 0)

    def h1(self, text: str) -> None:
        self.doc.add_heading(text, 1)

    def para(self, text: str, bold: bool = False, italic: bool = False) -> None:
        if not text.strip():
            raise ValueError("empty paragraphs are not allowed")
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic

    def label(self, label: str, text: str) -> None:
        """``Label : text`` with a bold label."""
        paragraph = self.doc.add_paragraph()
        paragraph.add_run(f"{label} :").bold = True
        paragraph.add_run(f" {text}")

    def bullets(self, items: Sequence[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def table(
        self,
        header: Sequence[str],
        rows: Sequence[Sequence[object]],
        widths: Sequence[float],
        right: Sequence[int] = (),
        bold_rows: Sequence[int] = (),
    ) -> None:
        """A grid table; ``right`` lists right-aligned columns, ``bold_rows`` 1-based data rows."""
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, values in enumerate([header, *rows]):
            for col, value in enumerate(values):
                cell = table.rows[r].cells[col]
                cell.text = str(value)
                paragraph = cell.paragraphs[0]
                if col in right:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.bold = r == 0 or r in bold_rows
        for row in table.rows:
            for col, width in enumerate(widths):
                row.cells[col].width = Cm(width)

    def render(self, w: World, doc_id: str) -> bytes:
        doc = w.doc(doc_id)
        meta = meta_for(doc, w.org_name(str(doc["organisation"])))
        buf = io.BytesIO()
        self.doc.save(buf)
        return finalize_package(buf.getvalue(), meta)


def _parties(b: DocxBuilder, w: World) -> None:
    parties = w.facts["contrat"]["parties"]
    b.para("Entre les soussignés :", bold=True)
    b.para(str(parties["entreprise"]))
    b.para("et")
    b.para(str(parties["client"]))
    b.para("Il a été convenu ce qui suit.")


def _signatures(b: DocxBuilder, w: World, keys: Sequence[str], day: str) -> None:
    builder = w.builder
    company = [k for k in keys if w.person(k)["organisation"] == builder["key"]]
    clients = [k for k in keys if w.person(k)["organisation"] == w.client["key"]]
    signed = f"Lu et approuvé, signé le {fr_date(day)}"
    b.table(
        ["Pour l'Entreprise", "Le Client"],
        [
            [
                "\n".join([builder["name"],
                           *(f"{w.name(k)}, {w.role(k).lower()}" for k in company), signed]),
                "\n".join([*(w.name(k) for k in clients), signed]),
            ]
        ],
        widths=(8.5, 8.5),
    )


def _articles(w: World) -> dict[str, str]:
    """Article titles by number (``"4"`` -> ``Article 4 - Échéancier de paiement``)."""
    titles: dict[str, str] = {}
    for raw in w.facts["contrat"]["articles"]:
        number, _, title = str(raw).partition(" ")
        titles[number] = f"Article {number} - {title}"
    return titles


# ------------------------------------------------------------------------------------------
# Works contract (signed 2026-02-10)


def contrat(w: World) -> bytes:
    check_contract(w)
    facts = w.facts["contrat"]
    devis = w.facts["devis"]
    site = w.place("chantier-moreau")
    builder = w.builder
    day = str(facts["date"])
    b = DocxBuilder(f"{builder['name']} - Contrat de travaux du {fr_date(day)}",
                    f"Contrat de travaux - {site['short']}, {w.town}",
                    legal=w.legal_line(builder["key"]))
    b.title(str(facts["title"]))
    _parties(b, w)
    article = _articles(w)
    b.h1(article["1"])
    b.para(_flat(facts["object"]))
    b.h1(article["2"])
    price = facts["price"]
    b.para(f"Le prix des travaux est fixé à {eur(price['ht'])} HT, soit {eur(price['ttc'])} TTC, "
           f"aux taux de TVA indiqués au devis {devis['number']} indice B. "
           f"Il s'agit d'un {price['nature']}.")
    b.h1(article["3"])
    b.para(f"Les travaux débutent le {fr_date(str(facts['start']))}. Leur durée est de "
           f"{facts['duration_weeks']} semaines ; l'achèvement des travaux est prévu le "
           f"{fr_date(str(facts['planned_completion']))}.")
    b.h1(article["4"])
    b.para("Le prix TTC est payé selon l'échéancier suivant :")
    schedule = facts["payment_schedule"]
    b.table(
        ["N°", "Échéance", "Pourcentage", "Montant TTC"],
        [
            *([str(item["n"]), _sentence(str(item["event"])), percent(item["percent"]),
               eur(item["amount"])] for item in schedule),
            ["", "Total", percent(sum(int(i["percent"]) for i in schedule)), eur(price["ttc"])],
        ],
        widths=(1.2, 10.3, 2.5, 3.0),
        right=(2, 3),
        bold_rows=(len(schedule) + 1,),
    )
    deposit = facts["deposit"]
    b.para(f"L'acompte de {percent(deposit['percent'])}, soit {eur(deposit['amount'])}, est versé "
           f"à la signature du présent contrat, le {fr_date(str(deposit['paid_on']))}.")
    b.h1(article["5"])
    retention = facts["retention"]
    b.para(f"Une retenue de garantie de {percent(retention['percent'])} du prix TTC, soit "
           f"{eur(retention['amount'])}, est appliquée. Elle est libérée {retention['release']} "
           "en l'absence de réserves non levées.")
    b.h1(article["6"])
    b.para(str(facts["subcontracting"]))
    b.h1(article["7"])
    b.para(str(facts["planning_permission"]))
    b.h1(article["8"])
    b.para(str(facts["insurance"]))
    b.h1(article["9"])
    b.para(str(facts["changes_clause"]))
    b.para(f"Fait à {facts['place_signed']}, le {fr_date(day)}, en deux exemplaires originaux.")
    b.label("Annexe", f"devis {devis['number']} indice B du "
                      f"{fr_date(str(devis['indice_b']['date']))}.")
    _signatures(b, w, facts["signatories"], day)
    return b.render(w, "contrat")


# ------------------------------------------------------------------------------------------
# Amendment n°1 (signed 2026-03-10): amends the contract, never supersedes it


def avenant(w: World) -> bytes:
    check_amendment(w)
    facts = w.facts["avenant"]
    contract = w.facts["contrat"]
    devis = w.facts["devis"]
    site = w.place("chantier-moreau")
    builder = w.builder
    day = str(facts["date"])
    b = DocxBuilder(f"{builder['name']} - {facts['title']}",
                    f"Avenant n°{facts['number']} - {site['short']}, {w.town}",
                    legal=w.legal_line(builder["key"]))
    b.title(str(facts["title"]))
    _parties(b, w)
    b.h1("Article 1 - Objet de l'avenant")
    b.para(_flat(facts["object"]))
    b.h1("Article 2 - Travaux supplémentaires")
    b.para("Les travaux suivants sont ajoutés aux travaux du contrat :")
    totals = facts["totals"]
    lines = facts["lines"]
    b.table(
        ["Réf.", "Désignation", "Qté", "U", "PU HT", "Total HT"],
        [
            *([str(line["ref"]), str(line["designation"]), quantity(line["qty"]),
               str(line["unit"]), eur(line["unit_price"]), eur(line["total"])] for line in lines),
            ["", "Total HT de l'avenant", "", "", "", eur(totals["ht"])],
            ["", f"TVA {percent(totals['vat_rate'])}", "", "", "", eur(totals["vat"])],
            ["", "Total TTC de l'avenant", "", "", "", eur(totals["ttc"])],
        ],
        widths=(1.2, 8.6, 1.2, 1.2, 2.4, 2.4),
        right=(2, 4, 5),
        bold_rows=(len(lines) + 1, len(lines) + 3),
    )
    b.h1("Article 3 - Nouveau montant du contrat")
    initial, new = facts["initial_contract"], facts["new_contract_total"]
    b.table(
        ["Montant", "HT", "TTC"],
        [
            [f"Contrat de travaux du {fr_date(str(contract['date']))} (devis {devis['number']} "
             "indice B)", eur(initial["ht"]), eur(initial["ttc"])],
            [f"Avenant n°{facts['number']}", eur(totals["ht"]), eur(totals["ttc"])],
            ["Nouveau montant du contrat", eur(new["ht"]), eur(new["ttc"])],
        ],
        widths=(10.0, 3.5, 3.5),
        right=(1, 2),
        bold_rows=(3,),
    )
    b.para(f"Le montant du contrat est ainsi porté à {eur(new['ht'])} HT, soit {eur(new['ttc'])} "
           "TTC.")
    b.h1("Article 4 - Paiement")
    b.para(str(facts["payment"]))
    b.h1("Article 5 - Délais")
    b.para(str(facts["schedule"]))
    b.para(f"Fait à {contract['place_signed']}, le {fr_date(day)}, en deux exemplaires originaux.")
    _signatures(b, w, facts["signatories"], day)
    return b.render(w, "avenant")


# ------------------------------------------------------------------------------------------
# Site meeting minutes n°3 (2026-03-24)


def compte_rendu(w: World) -> bytes:
    facts = w.facts["compte_rendu_3"]
    builder = w.builder
    site = w.place("chantier-moreau")
    day = str(facts["date"])
    title = f"Compte rendu de chantier n°{facts['number']}"
    b = DocxBuilder(f"{builder['name']} - {title}", f"{title} - {site['short']}, {w.town}")
    b.title(title)
    b.label("Date de la réunion", f"{fr_date(day)} à {facts['time']}")
    b.label("Lieu", str(facts["place"]))
    writer = facts["written_by"]
    b.label("Rédigé par", f"{w.name(writer)}, {w.role(writer).lower()} "
                          f"({w.org_name(w.person(writer)['organisation'])})")
    present, excused = list(facts["present"]), list(facts["excused"])
    people = [*present, *excused, *facts["distribution_extra"]]
    b.table(
        ["Nom", "Organisation", "Fonction", "Présent", "Excusé", "Diffusion"],
        [
            [w.name(key), w.org_name(w.person(key)["organisation"]), w.role(key),
             "x" if key in present else "", "x" if key in excused else "", "x"]
            for key in people
        ],
        widths=(3.6, 3.8, 4.4, 1.7, 1.7, 1.8),
    )
    b.h1("1. Avancement des travaux")
    b.bullets([str(item) for item in facts["progress"]])
    b.h1("2. Menuiseries extérieures")
    b.para(_flat(facts["windows"]))
    b.h1("3. Décisions")
    for decision in facts["decisions"]:
        b.para(f"{decision['id']} - {decision['text']}")
    b.h1("4. Actions")
    b.table(
        ["Action", "Responsable", "Échéance"],
        [
            [str(action["what"]), w.label(str(action["who"])), fr_date(str(action["due"]))]
            for action in facts["actions"]
        ],
        widths=(10.5, 3.8, 2.7),
    )
    b.h1("5. Prochaine réunion")
    meeting = facts["next_meeting"]
    b.para(f"Réunion de chantier n°{meeting['number']} le {fr_date(str(meeting['date']))} à "
           f"{meeting['time']}, sur le chantier.")
    return b.render(w, "cr3")
