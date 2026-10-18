"""Read generated sample files back: text per locator, counts and metadata.

Used by the manifest builder, the corpus tests and the golden-set tooling. Locator keys follow
docs/spec/evidence-contract.md section 2 (header = row 1 for XLSX and CSV).
"""

from __future__ import annotations

import csv
import email
import io
import re
import unicodedata
from dataclasses import dataclass, field
from email.policy import default as default_policy
from pathlib import Path
from typing import Any

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "csv": "text/csv",
    "eml": "message/rfc822",
    "vtt": "text/vtt",
    "md": "text/markdown",
    "txt": "text/plain",
    "jpg": "image/jpeg",
    "png": "image/png",
}

_VOICE = re.compile(r"<v\s+([^>]+)>")
_TIMING = re.compile(r"^(\d{2}):(\d{2}):(\d{2})\.(\d{3})\s+-->\s+(\d{2}):(\d{2}):(\d{2})\.(\d{3})")


@dataclass
class Extracted:
    """Text units of one file in reading order, keyed by single-unit locator keys."""

    fmt: str
    units: list[tuple[str, str]] = field(default_factory=list)
    furniture: list[str] = field(default_factory=list)
    counts: dict[str, Any] = field(default_factory=dict)
    metadata: dict[str, str] = field(default_factory=dict)

    def text(self) -> str:
        return "\n".join([*(t for _, t in self.units), *self.furniture])


def slug(value: str) -> str:
    """ASCII kebab-case slug (same rule as source slugs in the wiki contract)."""
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]", "-", ascii_value.lower())).strip("-")[:60]


def normalise(text: str) -> str:
    """Normalisation used for quote checks (compatibility form, case folding, quotes, spaces)."""
    text = unicodedata.normalize("NFKC", text).casefold()
    text = text.replace("\u00ad", "").replace("\u2019", "'").replace("\u2018", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    return re.sub(r"\s+", " ", text).strip()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _pdf(path: Path) -> Extracted:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(path))
    try:
        units = []
        for index in range(len(doc)):
            page = doc[index]
            textpage = page.get_textpage()
            units.append((f"p{index + 1:04d}", textpage.get_text_range()))
            textpage.close()
            page.close()
        metadata = {k: str(v) for k, v in doc.get_metadata_dict().items() if v}
        return Extracted("pdf", units, counts={"pages": len(units)}, metadata=metadata)
    finally:
        doc.close()


def _docx(path: Path) -> Extracted:
    from docx import Document

    doc = Document(str(path))
    units = [(f"para-{i:04d}", p.text) for i, p in enumerate(doc.paragraphs, start=1)]
    table_rows = []
    for t_index, table in enumerate(doc.tables, start=1):
        table_rows.append(len(table.rows))
        for r_index, row in enumerate(table.rows, start=1):
            key = f"tbl-{t_index:02d}-r{r_index:04d}-r{r_index:04d}"
            units.append((key, " | ".join(cell.text for cell in row.cells)))
    furniture = []
    for section in doc.sections:
        furniture += [p.text for p in section.header.paragraphs if p.text]
        furniture += [p.text for p in section.footer.paragraphs if p.text]
    props = doc.core_properties
    metadata = {
        "author": props.author or "",
        "last_modified_by": props.last_modified_by or "",
        "created": props.created.isoformat() if props.created else "",
        "modified": props.modified.isoformat() if props.modified else "",
        "title": props.title or "",
    }
    counts = {"paragraphs": len(doc.paragraphs), "tables": table_rows}
    return Extracted("docx", units, furniture, counts, metadata)


def _xlsx(path: Path) -> Extracted:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    units = []
    sheets = []
    furniture = []
    for ws in wb.worksheets:
        sheet_slug = slug(ws.title)
        rows = list(ws.iter_rows(values_only=True))
        sheets.append({"name": ws.title, "slug": sheet_slug, "rows": len(rows),
                       "columns": ws.max_column})
        for r_index, row in enumerate(rows, start=1):
            key = f"sheet-{sheet_slug}-r{r_index:04d}-r{r_index:04d}"
            units.append((key, " | ".join(_cell(v) for v in row)))
        for part in (ws.oddHeader, ws.oddFooter):
            for piece in (part.left, part.center, part.right):
                if piece.text:
                    furniture.append(piece.text)
    props = wb.properties
    metadata = {
        "author": props.creator or "",
        "last_modified_by": props.lastModifiedBy or "",
        "created": props.created.isoformat() if props.created else "",
        "modified": props.modified.isoformat() if props.modified else "",
        "title": props.title or "",
    }
    return Extracted("xlsx", units, furniture, {"sheets": sheets}, metadata)


def read_csv_rows(path: Path) -> tuple[list[list[str]], str]:
    raw = path.read_bytes().decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(raw.splitlines()[0], delimiters=",;")
    rows = list(csv.reader(io.StringIO(raw), delimiter=dialect.delimiter))
    return rows, dialect.delimiter


def _csv(path: Path) -> Extracted:
    rows, delimiter = read_csv_rows(path)
    units = [(f"r{i:04d}-r{i:04d}", " | ".join(row)) for i, row in enumerate(rows, start=1)]
    counts = {"rows": len(rows), "columns": len(rows[0]) if rows else 0, "delimiter": delimiter}
    return Extracted("csv", units, counts=counts)


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation

    prs = Presentation(str(path))
    units = []
    for index, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            parts.append(slide.notes_slide.notes_text_frame.text)
        units.append((f"slide-{index:02d}", "\n".join(parts)))
    props = prs.core_properties
    metadata = {
        "author": props.author or "",
        "last_modified_by": props.last_modified_by or "",
        "created": props.created.isoformat() if props.created else "",
        "modified": props.modified.isoformat() if props.modified else "",
        "title": props.title or "",
    }
    return Extracted("pptx", units, counts={"slides": len(units)}, metadata=metadata)


def _eml(path: Path) -> Extracted:
    msg = email.message_from_bytes(path.read_bytes(), policy=default_policy)
    headers = [f"{name}: {value}" for name, value in msg.items()]
    body_part = msg.get_body(preferencelist=("plain",))
    body = body_part.get_content() if body_part is not None else ""  # type: ignore[union-attr]
    lines = body.split("\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]
    units = [("msg-headers", "\n".join(headers))]
    units += [(f"msg-body-l{i:04d}-l{i:04d}", line) for i, line in enumerate(lines, start=1)]
    counts = {"body_lines": len(lines), "parts": len(list(msg.walk()))}
    return Extracted("eml", units, counts=counts)


def parse_vtt(text: str) -> list[tuple[str, str, str]]:
    """Return (start, end, text) for every cue; voice tags become 'Speaker: '."""
    cues = []
    for block in text.replace("\r\n", "\n").split("\n\n"):
        lines = [line for line in block.split("\n") if line.strip()]
        timing_index = next((i for i, line in enumerate(lines) if _TIMING.match(line)), None)
        if timing_index is None:
            continue
        match = _TIMING.match(lines[timing_index])
        assert match is not None
        start = f"{match.group(1)}:{match.group(2)}:{match.group(3)}"
        end = f"{match.group(5)}:{match.group(6)}:{match.group(7)}"
        body = " ".join(lines[timing_index + 1 :])
        body = _VOICE.sub(lambda m: f"{m.group(1)}: ", body)
        cues.append((start, end, body))
    return cues


def _vtt(path: Path) -> Extracted:
    cues = parse_vtt(path.read_text(encoding="utf-8"))
    units = [(f"cue-{start.replace(':', '-')}", text) for start, _, text in cues]
    return Extracted("vtt", units, counts={"cues": len(cues),
                                           "cue_starts": [start for start, _, _ in cues]})


def _lines(path: Path, fmt: str) -> Extracted:
    lines = path.read_text(encoding="utf-8").split("\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]
    units = [(f"l{i:04d}-l{i:04d}", line) for i, line in enumerate(lines, start=1)]
    return Extracted(fmt, units, counts={"lines": len(lines)})


def _image(path: Path, fmt: str) -> Extracted:
    from PIL import Image

    with Image.open(path) as img:
        width, height = img.size
    return Extracted(fmt, [], counts={"width": width, "height": height})


def extract(path: Path) -> Extracted:
    fmt = path.suffix.lower().lstrip(".")
    if fmt == "pdf":
        return _pdf(path)
    if fmt == "docx":
        return _docx(path)
    if fmt == "xlsx":
        return _xlsx(path)
    if fmt == "csv":
        return _csv(path)
    if fmt == "pptx":
        return _pptx(path)
    if fmt == "eml":
        return _eml(path)
    if fmt == "vtt":
        return _vtt(path)
    if fmt in ("md", "txt"):
        return _lines(path, fmt)
    if fmt in ("jpg", "png"):
        return _image(path, fmt)
    raise ValueError(f"unsupported format: {path}")


# ------------------------------------------------------------------------------------------
# Locator keys with ranges


_RANGE_PATTERNS = [
    (re.compile(r"^(sheet-[a-z0-9-]+?-)r(\d{4})-r(\d{4})$"), "r"),
    (re.compile(r"^(tbl-\d{2}-)r(\d{4})-r(\d{4})$"), "r"),
    (re.compile(r"^()r(\d{4})-r(\d{4})$"), "r"),
    (re.compile(r"^(msg-body-)l(\d{4})-l(\d{4})$"), "l"),
    (re.compile(r"^()l(\d{4})-l(\d{4})$"), "l"),
]


def expand_locator(key: str) -> list[str]:
    """Expand a range locator key into the single-unit keys produced by ``extract``."""
    for pattern, letter in _RANGE_PATTERNS:
        match = pattern.match(key)
        if match:
            prefix, start, end = match.group(1), int(match.group(2)), int(match.group(3))
            if end < start:
                raise ValueError(f"inverted range in locator {key}")
            return [f"{prefix}{letter}{i:04d}-{letter}{i:04d}" for i in range(start, end + 1)]
    return [key]


def locator_text(extracted: Extracted, key: str) -> str | None:
    """Text at a locator, or None when the locator does not exist (or is image/OCR-only).

    ``*`` is the golden-set wildcard for "any OCR region of an image": region numbering depends
    on the OCR engine, so image evidence has no deterministic text here.
    """
    if key == "*" or key.startswith("img-"):
        return None
    wanted = expand_locator(key)
    table = dict(extracted.units)
    if not all(k in table for k in wanted):
        return None
    return "\n".join(table[k] for k in wanted)


def find_units(extracted: Extracted, needle: str) -> list[str]:
    """Single-unit locator keys whose text contains ``needle`` (normalised comparison)."""
    target = normalise(needle)
    return [key for key, text in extracted.units if target in normalise(text)]
