"""PPTX training decks: onboarding and site rules, gearbox fundamentals, single-point lesson."""

from __future__ import annotations

import io
from collections.abc import Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt

from .common import STAMP, meta_for
from .diagrams import drivetrain_schematic, gearbox_exploded, sight_glass, to_png
from .ooxml import finalize_package
from .world import World

_BLUE = RGBColor(0x1F, 0x4E, 0x79)
_RED = RGBColor(0xC0, 0x00, 0x00)
_BAND = RGBColor(0xE9, 0xEF, 0xF7)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class DeckBuilder:
    def __init__(self, footer: str) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.footer = footer

    def _stamp(self, slide: Slide) -> None:
        box = slide.shapes.add_textbox(Inches(0.4), Inches(6.95), Inches(12.5), Inches(0.4))
        para = box.text_frame.paragraphs[0]
        para.text = f"{STAMP}  |  {self.footer}"
        para.runs[0].font.size = Pt(11)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = _RED

    def title_slide(self, title: str, subtitle: str) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle
        self._stamp(slide)

    def bullets(self, title: str, items: Sequence[str], notes: str = "",
                picture: bytes | None = None) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.left, body.top = Inches(0.6), Inches(1.5)
        body.width = Inches(7.4 if picture else 12.0)
        body.height = Inches(5.2)
        frame = body.text_frame
        frame.text = items[0]
        for item in items[1:]:
            frame.add_paragraph().text = item
        for para in frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(20)
        if picture:
            slide.shapes.add_picture(io.BytesIO(picture), Inches(8.3), Inches(1.7),
                                     width=Inches(4.6))
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self._stamp(slide)

    def table(self, title: str, header: Sequence[str], rows: Sequence[Sequence[str]],
              widths: Sequence[float], notes: str = "") -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        slide.shapes.title.text = title
        shape = slide.shapes.add_table(1 + len(rows), len(header), Inches(0.6), Inches(1.5),
                                       Inches(sum(widths)), Inches(0.45 * (1 + len(rows))))
        table = shape.table
        tbl_pr = table._tbl.tblPr
        for style_id in tbl_pr.findall(qn("a:tableStyleId")):
            tbl_pr.remove(style_id)
        for col, width in enumerate(widths):
            table.columns[col].width = Emu(int(Inches(width)))
        for col, value in enumerate(header):
            table.cell(0, col).text = value
        for r, values in enumerate(rows, start=1):
            for col, value in enumerate(values):
                table.cell(r, col).text = value
        for r in range(1 + len(rows)):
            for col in range(len(header)):
                cell = table.cell(r, col)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _BLUE if r == 0 else (_BAND if r % 2 else _WHITE)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(15)
                        if r == 0:
                            run.font.bold = True
                            run.font.color.rgb = _WHITE
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self._stamp(slide)

    def picture(self, title: str, png: bytes, caption: str, notes: str = "") -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(io.BytesIO(png), Inches(1.4), Inches(1.45), width=Inches(10.5))
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.4), Inches(12.0), Inches(0.5))
        box.text_frame.paragraphs[0].text = caption
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = _BLUE
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self._stamp(slide)

    def render(self, w: World, doc_id: str) -> bytes:
        doc = w.doc(doc_id)
        meta = meta_for(doc, w.org_name(str(doc["organisation"])))
        buf = io.BytesIO()
        self.prs.save(buf)
        return finalize_package(buf.getvalue(), meta)


def onboarding(w: World) -> bytes:
    d = DeckBuilder("Quellmoor Ridge onboarding and site rules - February 2025")
    d.title_slide("Quellmoor Ridge Wind Farm - onboarding and site rules",
                  "Contoso Renewables - induction for new staff and contractors - February 2025")
    d.bullets(
        "Welcome to Quellmoor Ridge",
        [
            "12 Litware LT-32 turbines in CCV configuration, WTG-01 to WTG-12, 38.4 MW",
            "Commercial operation since 15 October 2019",
            "Hub height 92 m, rotor diameter 118 m",
            "Grid connection to City Power & Light at the Skelvane 132 kV substation",
            "O&M building with workshop and stores at the site entrance",
        ],
        notes="Show the site layout plan on the wall of the meeting room.",
    )
    d.table(
        "Who is who",
        ["Role", "Name"],
        [
            ["Site Manager", "Maren Ellingsby"],
            ["O&M Supervisor", "Dario Penwarden"],
            ["HSE Coordinator", "Greta Vollan"],
            ["Reliability Engineer", "Yusuf Arkwell"],
            ["Stores and spare parts", "Tomasz Lebrecht"],
            ["Lead Wind Technician", "Callum Rydings"],
        ],
        widths=(5.0, 5.0),
    )
    d.bullets(
        "Site rules",
        [
            "Speed limit 20 km/h on all site roads",
            "Sign in and out at the O&M building",
            "Two persons minimum for any turbine access - no lone working",
            "Helmet, safety boots and high-visibility vest outside vehicles",
            "Report every incident and near miss on the same day",
        ],
    )
    d.bullets(
        "Permits and isolation",
        [
            "Every job on a turbine needs a PTW signed by the issuer",
            "Isolation follows SOP-LOTO-001 with a personal padlock and name tag",
            "High-voltage switching only by authorised persons",
            "The control room must always know who is in which turbine",
        ],
    )
    d.bullets(
        "Climbing and rescue",
        [
            "Follow SOP-WAH-003 for all work at height",
            "Climbing allowed up to a 10-minute mean wind speed of 15 m/s; hub entry up to 12 m/s",
            "A sealed rescue kit is kept in every nacelle - check the seal every day",
            "Rescue target: casualty on the ground within 15 minutes",
        ],
    )
    d.bullets(
        "Weather and ice",
        [
            "Ice detection (alarm A7001) stops the turbine automatically",
            "Keep 300 m away from a turbine with ice on the blades",
            "No climbing when lightning is detected within 15 km",
            "In cold weather work in pairs and take warm breaks in the vehicle",
        ],
    )
    d.bullets(
        "Emergency",
        [
            "Raise the alarm to the control room first",
            "Muster point: O&M building car park",
            "First aiders on every shift; defibrillator in the O&M building and in each service "
            "vehicle",
            "Follow the site ERP",
        ],
    )
    d.table(
        "Acronyms we use",
        ["Acronym", "Meaning"],
        [
            ["CCV", "Cold Climate Version - LT-32 variant with gearbox oil heaters, "
                    "low-temperature steel and heated sensors"],
            ["PTW", "Permit to work"],
            ["LOTO", "Lockout/tagout"],
            ["MOC", "Management of change"],
            ["TBT", "Toolbox talk"],
            ["SPL", "Single-point lesson"],
            ["WTG", "Wind turbine generator (our turbine numbers WTG-01 to WTG-12)"],
            ["ERP", "Emergency response plan"],
            ["CMS", "Condition monitoring system"],
            ["SB", "Service bulletin (issued by the turbine manufacturer)"],
        ],
        widths=(2.0, 10.0),
        notes="Ask the participants which acronyms they already know.",
    )
    d.bullets(
        "Environment",
        [
            "The site has an environmental permit with conditions on noise, shadow flicker, ice "
            "and wildlife",
            "Some turbines stop automatically for permit reasons - never override these stops",
            "A spill kit is kept in every turbine",
            "Report any spill immediately to the HSE Coordinator",
        ],
    )
    d.bullets(
        "Questions?",
        [
            "HSE Coordinator: Greta Vollan",
            "E-mail: greta.vollan@contoso-renewables.example",
            "Keep this induction record with your training file",
        ],
    )
    return d.render(w, "deck-onboarding")


def gearbox_fundamentals(w: World) -> bytes:
    d = DeckBuilder("Gearbox fundamentals - October 2025")
    d.title_slide("Gearbox fundamentals for wind technicians",
                  "Fabrikam FGT-3200 in the Litware LT-32 - Yusuf Arkwell - October 2025")
    d.picture("From rotor to generator", to_png(drivetrain_schematic(), dpi=110),
              "The gearbox raises 12.7 rpm at the rotor to 1,500 rpm at the generator.")
    d.bullets(
        "Three stages",
        [
            "Planetary stage: planet carrier, three planet gears, ring gear and sun pinion",
            "Two parallel helical stages",
            "Overall ratio 1:118.4",
            "Rotor 12.7 rpm, generator 1,500 rpm",
        ],
    )
    d.table(
        "Shaft names and acronyms",
        ["Acronym", "Meaning"],
        [
            ["LSS", "Low-speed shaft - main shaft on the rotor side"],
            ["IMS", "Intermediate-speed shaft - between the two helical stages"],
            ["HSS", "High-speed shaft - output shaft to the generator coupling"],
            ["PQ index", "Particle quantifier index - amount of ferrous debris in an oil sample"],
            ["RMS", "Root mean square - value used for vibration velocity"],
        ],
        widths=(2.4, 9.6),
    )
    d.bullets(
        "Lubrication: the change to PWG 220",
        [
            "Since SB-LWS-2025-07 our gearboxes use Proseware PWG 220 (ISO VG 220) instead of "
            "PWG 320",
            "410 L per gearbox; 60 L flush during the conversion",
            "Minimum oil temperature for start-up: +5 °C (it was +10 °C with PWG 320)",
            "Never mix oils: top up with PWG 220 only",
        ],
        notes="The conversion campaign runs from September to November 2025 (MOC-2025-012).",
    )
    d.bullets(
        "Temperatures and alarms",
        [
            "Normal oil sump temperature 50 to 65 °C",
            "A2101 warning at 75 °C; A2102 trip at 82 °C",
            "The thermostatic valve opens the cooler circuit at 45 °C",
            "Keep the cooler fan and the radiator clean",
        ],
    )
    d.bullets(
        "Condition monitoring",
        [
            "Vibration velocity RMS on the HSS bearings: alert 3.5 mm/s, alarm 4.5 mm/s",
            "Oil analysis every 6 months: iron alert 50 ppm, PQ index alert 60",
            "The trend matters more than a single value",
        ],
    )
    d.bullets(
        "Typical failure modes",
        [
            "HSS bearing damage: spalling and axial cracks",
            "Micropitting on gear teeth",
            "Cooler fouling and fan motor failures",
            "Oil leaks at seals and hoses",
        ],
    )
    d.picture("Exploded view of the FGT-3200", to_png(gearbox_exploded(), dpi=110),
              "Item numbers as in the maintenance manual LWS-MM-GBX-32, figure 4-1.")
    d.bullets(
        "Quiz",
        [
            "1. Which shaft drives the generator?",
            "2. At which oil temperature does A2102 trip the turbine?",
            "3. Which oil do we use since the conversion?",
            "4. What does a rising PQ index tell you?",
        ],
        notes="Answers: 1. the HSS. 2. 82 °C. 3. Proseware PWG 220. 4. More ferrous wear debris, "
              "often from a bearing.",
    )
    return d.render(w, "deck-gearbox")


def spl_oil_level(w: World) -> bytes:
    d = DeckBuilder("SPL-012 Gearbox oil level check - 2025-09-29")
    d.title_slide("SPL-012 - Gearbox oil level check",
                  "Single-point lesson - Quellmoor Ridge - 2025-09-29 - Callum Rydings")
    d.bullets(
        "How to check",
        [
            "Stop the turbine and wait 20 minutes so the oil settles",
            "Read the sight glass on the generator side of the gearbox",
            "The level must be between MIN and MAX",
            "Top up only with PWG 220 through the filter cart (10 micron) - never PWG 320",
            "Record the result on the daily checklist (item 4)",
        ],
        picture=to_png(sight_glass(), dpi=110),
    )
    d.picture("Good or not good?", to_png(sight_glass(), dpi=110),
              "Level below MIN: do not restart, report to the O&M Supervisor (risk of alarm "
              "A2130).")
    return d.render(w, "deck-spl")
